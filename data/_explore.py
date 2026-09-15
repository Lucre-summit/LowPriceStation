import openpyxl, re, collections, json

wb = openpyxl.load_workbook('data/raw/PEA_VOLTA_May_2024.xlsx', data_only=True)
ws = wb['PEA VOLTA 409 (เปิดแล้ว)']
hdr = list(next(ws.iter_rows(min_row=4, max_row=4, values_only=True)))
rows = [r for r in ws.iter_rows(min_row=5, max_row=415, values_only=True)]
H = {h: i for i, h in enumerate(hdr) if h}
print('rows', len(rows), 'seq-none', sum(1 for r in rows if r[0] is None))

cnt = collections.Counter()
headtexts = collections.Counter()
cabtexts = collections.Counter()
for r in rows:
    cab = str(r[H['ตู้ชาร์จ']])
    h1 = str(r[H['หัวชาร์จเครื่องที่ 1']] or '')
    h2 = str(r[H['หัวชาร์จเครื่องที่ 2']] or '')
    bands = [(int(m.group(1)), int(m.group(2))) for m in re.finditer(r'(\d+)kW\s*(\d+)\s*เครื่อง', cab)]
    groups = [g for g in (h1, h2) if g and g != '-']
    cnt[(sum(c for _, c in bands), len(groups))] += 1
    cabtexts[cab] += 1
    for g in groups:
        headtexts[g] += 1
print('(cabinets, headgroups):', sorted(cnt.items()))
print('--- head texts')
for k, v in headtexts.most_common(30):
    print(v, '|', k)
print('--- cabinet texts (top 25)')
for k, v in cabtexts.most_common(25):
    print(v, '|', k)
print('--- provinces')
print(collections.Counter(r[H['จังหวัด']] for r in rows).most_common(15))
print('--- khet values for metro')
metro = {'กทม.', 'นนทบุรี', 'ปทุมธานี', 'สมุทรปราการ', 'สมุทรสาคร', 'นครปฐม'}
sub = [r for r in rows if r[H['จังหวัด']] in metro]
print(len(sub))
for r in sub:
    print(r[0], r[H['ชื่อสถานี PEA ภาษาไทย']], '|', r[H['จังหวัด']], r[H['อำเภอ']], '|', r[H['ละติดจูด']], r[H['ลองติดจูด']], '|',
          r[H['25kW']], r[H['50kW']], r[H['120kW']], r[H['300kW']], r[H['360kW']], '| T2', r[H['AC Type 2']], 'CCS2', r[H['CCS2']], 'CHA', r[H['CHAdeMO']], '|', str(r[H['หัวชาร์จเครื่องที่ 1']])[:40], '/', str(r[H['หัวชาร์จเครื่องที่ 2']])[:40])
