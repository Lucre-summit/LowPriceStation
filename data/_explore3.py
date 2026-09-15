import openpyxl, re, collections, json, math

wb = openpyxl.load_workbook('data/raw/PEA_VOLTA_May_2024.xlsx', data_only=True)
ws = wb['PEA VOLTA 409 (เปิดแล้ว)']
hdr = list(next(ws.iter_rows(min_row=4, max_row=4, values_only=True)))
rows = [r for r in ws.iter_rows(min_row=5, max_row=415, values_only=True)]
H = {h: i for i, h in enumerate(hdr) if h}
BANDS = [25, 50, 120, 300, 360]

for r in rows:
    if not isinstance(r[0], (int, float)):
        print('NONNUM SEQ row:', r[:12])

for i, r in enumerate(rows):
    if r[H['ชื่อสถานี PEA ภาษาไทย']] == 'PEA VOLTA ปากช่อง':
        print('PAKCHONG row', i + 5, r[:12], r[H['ฐาน'] if 'ฐาน' in H else 0])

print('=== band column vs text consistency')
bad = 0
for r in rows:
    cab = str(r[H['ตู้ชาร์จ']])
    text = []
    for m in re.finditer(r'(\d+)kW\s*(\d+)\s*เครื่อง', cab):
        text += [int(m.group(1))] * int(m.group(2))
    cols = []
    for kw in BANDS:
        n = r[H[f'{kw}kW']]
        n = int(n) if n not in (None, '') else 0
        cols += [kw] * n
    if text != cols:
        bad += 1
        print('BANDDIFF', r[0], r[H['ชื่อสถานี PEA ภาษาไทย']], text, cols)
print('band diffs:', bad)

print('=== per-standard consistency: heads vs count columns')
bad = 0
for r in rows:
    groups = [g for g in (str(r[H['หัวชาร์จเครื่องที่ 1']] or ''), str(r[H['หัวชาร์จเครื่องที่ 2']] or '')) if g and g != '-']
    exp = collections.Counter({'AC Type 2': 0, 'CCS2': 0, 'CHAdeMO': 0})
    for g in groups:
        for part in g.strip().strip('()').split(','):
            exp[part.strip()] += 1
    for std, col in (('AC Type 2', 'AC Type 2'), ('CCS2', 'CCS2'), ('CHAdeMO', 'CHAdeMO')):
        v = r[H[col]]
        v = int(v) if v not in (None, '') else 0
        if exp[std] != v:
            bad += 1
            print('STDDIFF', r[0], r[H['ชื่อสถานี PEA ภาษาไทย']], std, 'heads', exp[std], 'col', v)
print('std diffs:', bad)

print('=== coords sanity')
for r in rows:
    try:
        lat = float(r[H['ละติดจูด']]); lng = float(r[H['ลองติดจูด']])
    except Exception:
        print('BADCOORD', r[0], r[H['ชื่อสถานี PEA ภาษาไทย']], r[H['ละติดจูด']], r[H['ลองติดจูด']])
        continue
    if not (5.5 < lat < 20.5 and 97 < lng < 106):
        print('OUTSIDE', r[0], lat, lng, r[H['ชื่อสถานี PEA ภาษาไทย']])

print('=== BMA arcgis bounds')
d = json.load(open('data/raw/bma_arcgis_dcevbmA_chargingstation.json', encoding='utf-8'))
lats = [f['attributes']['LATITUDE'] for f in d['features']]
lngs = [f['attributes']['LONGITUDE'] for f in d['features']]
print(min(lats), max(lats), min(lngs), max(lngs))
print('null lat/lng', sum(1 for x in lats if x is None), sum(1 for x in lngs if x is None))
gh = [f['geometry'] for f in d['features']]
print('geom none', sum(1 for g in gh if g is None))
print('=== POWER values')
for k, v in collections.Counter(f['attributes']['POWER'] for f in d['features']).most_common(40):
    print(v, '|', repr(k))
